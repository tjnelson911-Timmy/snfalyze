"""
Document utilities for upload processing, checksums, and classification
"""
import hashlib
import re
from typing import Optional, Tuple, Dict, Any


def calculate_checksum(content: bytes) -> str:
    """Calculate SHA-256 checksum of file content"""
    return hashlib.sha256(content).hexdigest()


# Document type classification patterns
DOC_TYPE_PATTERNS = {
    # Offering Memorandum / CIM
    "om": {
        "filename_patterns": [
            r"(?:offering|om|cim|teaser|marketing)",
            r"(?:confidential.*info.*memo)",
            r"(?:investment.*memo)",
        ],
        "content_patterns": [
            r"offering\s+memorandum",
            r"confidential\s+information\s+memorandum",
            r"investment\s+opportunity",
            r"executive\s+summary",
            r"investment\s+highlights",
            r"investment\s+thesis",
            r"price\s+guidance",
        ],
        "weight": 1.0
    },
    # Financial Statements
    "financials": {
        "filename_patterns": [
            r"(?:financial|p&l|p\s*&\s*l|income|pnl)",
            r"(?:statement|profit.*loss)",
            r"(?:trailing|t-?12|ttm)",
            r"(?:budget|forecast|pro\s*forma)",
        ],
        "content_patterns": [
            r"income\s+statement",
            r"profit\s+and\s+loss",
            r"operating\s+statement",
            r"revenue",
            r"expense",
            r"ebitda",
            r"net\s+income",
            r"total\s+revenue",
        ],
        "weight": 1.0
    },
    # General Ledger / Trial Balance
    "gl": {
        "filename_patterns": [
            r"(?:general\s*ledger|gl|trial\s*balance|tb)",
        ],
        "content_patterns": [
            r"general\s+ledger",
            r"trial\s+balance",
            r"account\s+number",
            r"debit",
            r"credit",
        ],
        "weight": 0.9
    },
    # Rent Roll
    "rent_roll": {
        "filename_patterns": [
            r"(?:rent\s*roll|lease\s*roll|tenant)",
        ],
        "content_patterns": [
            r"rent\s+roll",
            r"tenant",
            r"unit\s+mix",
            r"lease\s+term",
            r"monthly\s+rent",
        ],
        "weight": 1.0
    },
    # Census / Occupancy
    "census": {
        "filename_patterns": [
            r"(?:census|occupancy|patient\s*days)",
            r"(?:daily\s*census|admissions)",
        ],
        "content_patterns": [
            r"daily\s+census",
            r"patient\s+days",
            r"occupancy",
            r"admissions",
            r"discharges",
            r"average\s+daily\s+census",
        ],
        "weight": 1.0
    },
    # Payor Mix
    "payor_mix": {
        "filename_patterns": [
            r"(?:payor|payer|mix|reimbursement)",
        ],
        "content_patterns": [
            r"payor\s+mix",
            r"payer\s+mix",
            r"medicaid",
            r"medicare",
            r"private\s+pay",
            r"managed\s+care",
            r"reimbursement",
        ],
        "weight": 0.9
    },
    # Survey / Inspection
    "survey": {
        "filename_patterns": [
            r"(?:survey|inspection|deficiency|citation)",
            r"(?:cms\s*2567|state\s*survey)",
        ],
        "content_patterns": [
            r"survey",
            r"deficiency",
            r"citation",
            r"plan\s+of\s+correction",
            r"cms\s*2567",
            r"immediate\s+jeopardy",
            r"scope\s+and\s+severity",
        ],
        "weight": 1.0
    },
    # Quality / Star Ratings
    "quality": {
        "filename_patterns": [
            r"(?:quality|star\s*rating|qm|measures)",
        ],
        "content_patterns": [
            r"star\s+rating",
            r"quality\s+measure",
            r"overall\s+rating",
            r"health\s+inspection",
            r"staffing\s+rating",
        ],
        "weight": 0.8
    },
    # Lease
    "lease": {
        "filename_patterns": [
            r"(?:lease|rental\s*agreement|sublease)",
        ],
        "content_patterns": [
            r"lease\s+agreement",
            r"landlord",
            r"tenant",
            r"rental\s+payment",
            r"term\s+of\s+lease",
            r"base\s+rent",
            r"triple\s+net",
        ],
        "weight": 1.0
    },
    # License / Certification
    "license": {
        "filename_patterns": [
            r"(?:license|certification|permit|accreditation)",
        ],
        "content_patterns": [
            r"license",
            r"certification",
            r"certificate\s+of\s+need",
            r"licensed\s+beds",
            r"accreditation",
        ],
        "weight": 0.8
    },
    # Appraisal
    "appraisal": {
        "filename_patterns": [
            r"(?:appraisal|valuation|assessment)",
        ],
        "content_patterns": [
            r"appraisal",
            r"fair\s+market\s+value",
            r"comparable\s+sales",
            r"income\s+approach",
            r"cost\s+approach",
        ],
        "weight": 0.9
    },
    # Staffing
    "staffing": {
        "filename_patterns": [
            r"(?:staffing|personnel|employee|roster)",
        ],
        "content_patterns": [
            r"staffing",
            r"fte",
            r"full\s+time\s+equivalent",
            r"nursing\s+hours",
            r"hprd",
            r"rn\s+hours",
            r"cna\s+hours",
        ],
        "weight": 0.8
    },
    # Capital Expenditure
    "capex": {
        "filename_patterns": [
            r"(?:capex|capital|improvement|renovation)",
        ],
        "content_patterns": [
            r"capital\s+expenditure",
            r"renovation",
            r"improvement",
            r"construction",
            r"equipment",
        ],
        "weight": 0.7
    },
}


def classify_document_type(filename: str, content_sample: str = "") -> Tuple[str, float]:
    """
    Classify document type based on filename and content.
    Returns (doc_type, confidence).
    """
    filename_lower = filename.lower()
    content_lower = content_sample.lower() if content_sample else ""

    scores = {}

    for doc_type, patterns in DOC_TYPE_PATTERNS.items():
        score = 0.0
        matches = 0

        # Check filename patterns
        for pattern in patterns["filename_patterns"]:
            if re.search(pattern, filename_lower):
                score += 0.5
                matches += 1

        # Check content patterns (if content provided)
        if content_lower:
            for pattern in patterns["content_patterns"]:
                if re.search(pattern, content_lower):
                    score += 0.3
                    matches += 1

        # Apply weight
        if matches > 0:
            # More matches = higher confidence
            confidence = min(1.0, (score * patterns["weight"]) / (1.0 + 0.1 * matches))
            scores[doc_type] = confidence

    if not scores:
        return "other", 0.0

    # Return type with highest score
    best_type = max(scores, key=scores.get)
    return best_type, scores[best_type]


def get_doc_type_display_name(doc_type: str) -> str:
    """Get human-readable display name for document type"""
    names = {
        "om": "Offering Memorandum",
        "financials": "Financial Statement",
        "gl": "General Ledger",
        "rent_roll": "Rent Roll",
        "census": "Census/Occupancy",
        "payor_mix": "Payor Mix",
        "survey": "Survey/Inspection",
        "quality": "Quality Report",
        "lease": "Lease Agreement",
        "license": "License/Certification",
        "appraisal": "Appraisal",
        "staffing": "Staffing Report",
        "capex": "Capital Expenditure",
        "other": "Other Document",
    }
    return names.get(doc_type, doc_type.replace("_", " ").title())


def extract_text_sample(file_path: str, file_type: str, max_chars: int = 5000) -> str:
    """
    Extract a sample of text from document for classification.
    Uses quick extraction methods, not full parsing.
    """
    try:
        if file_type == "pdf":
            return _extract_pdf_sample(file_path, max_chars)
        elif file_type in ["docx", "doc"]:
            return _extract_docx_sample(file_path, max_chars)
        elif file_type in ["xlsx", "xls"]:
            return _extract_excel_sample(file_path, max_chars)
        elif file_type == "csv":
            return _extract_csv_sample(file_path, max_chars)
    except Exception as e:
        print(f"Error extracting text sample: {e}")
    return ""


def _extract_pdf_sample(file_path: str, max_chars: int) -> str:
    """Extract text sample from PDF"""
    try:
        import PyPDF2
        text = ""
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            # Just read first few pages
            for i, page in enumerate(reader.pages[:3]):
                page_text = page.extract_text() or ""
                text += page_text
                if len(text) >= max_chars:
                    break
        return text[:max_chars]
    except:
        return ""


def _extract_docx_sample(file_path: str, max_chars: int) -> str:
    """Extract text sample from Word document"""
    try:
        from docx import Document
        doc = Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs[:50]])
        return text[:max_chars]
    except:
        return ""


def _extract_excel_sample(file_path: str, max_chars: int) -> str:
    """Extract text sample from Excel file"""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        text_parts = []

        for sheet in wb.worksheets[:2]:  # First 2 sheets
            # Get sheet name
            text_parts.append(sheet.title)

            # Get first 20 rows
            for row_idx, row in enumerate(sheet.iter_rows(max_row=20, values_only=True)):
                if row_idx > 20:
                    break
                row_text = " ".join(str(cell) if cell else "" for cell in row)
                text_parts.append(row_text)

        text = "\n".join(text_parts)
        return text[:max_chars]
    except:
        return ""


def _extract_csv_sample(file_path: str, max_chars: int) -> str:
    """Extract text sample from CSV file"""
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read(max_chars)
    except:
        return ""
